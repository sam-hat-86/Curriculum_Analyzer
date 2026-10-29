"""
2ページ目以降 (page_index >= 2) のキー整合性・Excel出力JOIN検証テスト (C1再発防止)
"""
import openpyxl
from src.database.db_manager import DatabaseManager
from src.database.repository import Repository
from src.exporter.excel_exporter import ExcelExporter
from src.models.curriculum import (
    CurriculumOverview,
    TargetInfo,
    UnitRecord,
    TargetProgress,
    CurriculumDetailData,
)
from src.models.evaluation import InstructionEvaluation
from src.models.state import ProcessState, TimingCategory
from src.utils.constants import (
    SHEET_SUMMARY_RESULT,
    SHEET_PROGRESS,
    SHEET_UNIT_DETAILS,
    SHEET_INSTRUCTION,
)


def test_page2_and_page3_key_consistency_and_excel_join(tmp_path):
    """
    page_index=2, 3 のカリキュラムについて、
    save_pages_init -> update_page_fetched -> save_parsed_result を実行した際、
    1. pages.status が確実に SAVED になり、get_uncompleted_pages() に残らないこと
    2. 単元詳細シート・備考欄シート・集計結果シートの JOIN で欠落せず全件出力されること
    """
    db_file = str(tmp_path / "test_page_key.db")
    db_manager = DatabaseManager(db_file)
    repo = Repository(db_manager)

    ov_p1 = CurriculumOverview(
        classroom_name="梅田校", classroom_code="1001", school_year="2026",
        student_id="S001", student_name="一頁 太郎", grade="高1", division="通常",
        subject="英語", row_index=0, page_index=1
    )
    ov_p2 = CurriculumOverview(
        classroom_name="梅田校", classroom_code="1001", school_year="2026",
        student_id="S002", student_name="二頁 花子", grade="高2", division="通常",
        subject="数学", row_index=0, page_index=2
    )
    ov_p3 = CurriculumOverview(
        classroom_name="梅田校", classroom_code="1001", school_year="2026",
        student_id="S003", student_name="三頁 次郎", grade="高3", division="講習",
        subject="物理", row_index=5, page_index=3
    )

    assert ov_p1.page_key == "1001_S001_英語_0"
    assert ov_p2.page_key == "1001_S002_数学_2_0"
    assert ov_p3.page_key == "1001_S003_物理_3_5"

    repo.save_pages_init([ov_p1, ov_p2, ov_p3])
    assert len(repo.get_uncompleted_pages()) == 3

    for ov in (ov_p1, ov_p2, ov_p3):
        repo.update_page_fetched(ov.page_key, "<html><body>detail</body></html>", "2026-10-29 10:00:00")
        t = TargetInfo(
            target_no="T1", target_name="期末", start_date="2026/05/01", end_date="2026/06/30",
            period_str="2026/05/01～2026/06/30", is_current=True, target_score="80",
            current_score="75", site_progress_rate="50%", progress_delay="進捗遅れ"
        )
        u = UnitRecord(
            unit_no="1", unit_name=f"単元_{ov.student_name}", middle_unit="基礎",
            textbook="標準テキスト", problem_code="P1", score="8", max_score="10", score_rate="80",
            execution_date="2026/05/10", timing=TimingCategory.ON_TIME, is_executed=True,
            today_textbook="標準テキスト", execution_division="通常", confirm_test="合格",
            target_checks={"T1": True}
        )
        p = TargetProgress(
            classroom=ov.classroom_name, student_id=ov.student_id, grade=ov.grade,
            student_name=ov.student_name, division=ov.division, subject=ov.subject,
            target_no="T1", target_name="期末", start_date="2026/05/01", end_date="2026/06/30",
            period_str="2026/05/01～2026/06/30", is_current=True, textbooks_used="標準テキスト",
            target_score="80", current_score="75", target_unit_count=1, executed_count=1,
            executed_rate="100", on_time_count=1, on_time_rate="100", before_count=0,
            before_rate="0", after_count=0, after_rate="0", unexecuted_count=0,
            unexecuted_rate="0", latest_execution_date="2026/05/10", progress_delay="進捗遅れ"
        )
        ev = InstructionEvaluation(
            severity="WARNING", author="講師A", textbook="標準テキスト(所持)",
            warnings=["WARN_TEST_CRITERIA_MISSING"], fix_items=["小テスト基準を記載"]
        )
        detail = CurriculumDetailData(
            overview=ov,
            raw_html="<html><body>detail</body></html>",
            fetch_time="2026-10-29 10:00:00",
            targets=[t],
            units=[u],
            instruction_eval=ev,
            progress_list=[p],
        )
        repo.save_parsed_result(detail)

    counts = repo.get_counts()
    assert counts["total"] == 3
    assert counts["saved"] == 3
    assert counts["fetched"] == 0
    assert counts["unfetched"] == 0
    assert len(repo.get_uncompleted_pages()) == 0

    out_path = str(tmp_path / "page_consistency.xlsx")
    exporter = ExcelExporter(repo)
    exporter.export_to(out_path)

    wb = openpyxl.load_workbook(out_path)
    try:
        # 単元詳細シートに3名全員が出力されていること
        ws_unit = wb[SHEET_UNIT_DETAILS]
        unit_students = [ws_unit.cell(row=r, column=4).value for r in range(2, ws_unit.max_row + 1)]
        assert unit_students == ["一頁 太郎", "二頁 花子", "三頁 次郎"]

        # 備考欄シートに3名全員が出力されていること
        ws_inst = wb[SHEET_INSTRUCTION]
        inst_students = [ws_inst.cell(row=r, column=4).value for r in range(2, ws_inst.max_row + 1)]
        assert inst_students == ["一頁 太郎", "二頁 花子", "三頁 次郎"]

        # 集計結果シートに3名全員が出力されていること
        ws_res = wb[SHEET_SUMMARY_RESULT]
        res_students = [ws_res.cell(row=r, column=2).value for r in range(2, ws_res.max_row + 1)]
        assert res_students == ["一頁 太郎", "二頁 花子", "三頁 次郎"]
    finally:
        wb.close()

