"""
デモ生徒（「デモ 一（デモ イチ）」等）のスキップ検証テスト
"""
import openpyxl
from src.utils.normalization import is_demo_student
from src.parser.list_parser import ListParser
from src.database.db_manager import DatabaseManager
from src.database.repository import Repository
from src.exporter.excel_exporter import ExcelExporter
from src.models.curriculum import CurriculumOverview, TargetProgress, CurriculumDetailData
from src.models.evaluation import InstructionEvaluation
from src.analyzer.copypaste_detector import detect_copy_paste

def test_is_demo_student_matching():
    """デモ生徒判定の各種パターンと非該当パターンの検証"""
    # 1. ユーザー指定の「デモ 一（デモ イチ）」とその表記揺れ
    assert is_demo_student("デモ 一（デモ イチ）") is True
    assert is_demo_student("デモ 一(デモ イチ)") is True
    assert is_demo_student("デモ 一 (デモ イチ)") is True
    assert is_demo_student("デモ 一 (デモイチ)") is True
    assert is_demo_student("デモ一（デモイチ）") is True
    assert is_demo_student("デモ一(デモイチ)") is True
    assert is_demo_student("デモ　一（デモ　イチ）") is True # 全角空白

    # 2. フリガナなしの「デモ 一」
    assert is_demo_student("デモ 一") is True
    assert is_demo_student("デモ一") is True

    # 3. 名字が「デモ」のデモ生徒全般
    assert is_demo_student("デモ 太郎") is True
    assert is_demo_student("デモ 花子") is True
    assert is_demo_student("デモ") is True

    # 4. 「デモ田」「デモ川」など姓がデモ単体でない生徒 (除外してはならない)
    assert is_demo_student("デモ田 太郎") is False
    assert is_demo_student("デモ川 花子") is False
    assert is_demo_student("デモ山 一郎") is False

    # 5. 一般生徒
    assert is_demo_student("山田 太郎") is False
    assert is_demo_student("田中 二郎") is False
    assert is_demo_student("佐藤 健") is False

    # 6. 空値・境界値
    assert is_demo_student("") is False
    assert is_demo_student("   ") is False
    assert is_demo_student(None) is False


def test_list_parser_skips_demo_student():
    """一覧パース時にデモ生徒が除外されることを検証"""
    sample_html = """
    <table>
        <thead>
            <tr>
                <th>学校/コース</th>
                <th>学年</th>
                <th>受講区分</th>
                <th>科目</th>
                <th>生徒</th>
                <th>カリキュラム</th>
            </tr>
        </thead>
        <tbody>
            <tr>
                <td>高校A</td>
                <td>高3</td>
                <td>通常</td>
                <td>英語</td>
                <td>00000001 山田 太郎</td>
                <td><button>シミュレーションシート</button></td>
            </tr>
            <tr>
                <td>高校B</td>
                <td>高2</td>
                <td>通常</td>
                <td>数学</td>
                <td>00000002 デモ 一（デモ イチ）</td>
                <td><button>シミュレーションシート</button></td>
            </tr>
            <tr>
                <td>高校C</td>
                <td>高1</td>
                <td>通常</td>
                <td>国語</td>
                <td>00000003 佐藤 花子</td>
                <td><button>シミュレーションシート</button></td>
            </tr>
            <tr>
                <td>高校D</td>
                <td>中3</td>
                <td>通常</td>
                <td>理科</td>
                <td>00000004 デモ 太郎</td>
                <td><button>シミュレーションシート</button></td>
            </tr>
        </tbody>
    </table>
    """
    parser = ListParser()
    overviews = parser.parse(sample_html, classroom_name="テスト校")

    # 4行中、デモ生徒2行がスキップされ、一般生徒2行のみ残る
    assert len(overviews) == 2
    assert overviews[0].student_name == "山田 太郎"
    assert overviews[0].student_id == "00000001"
    assert overviews[1].student_name == "佐藤 花子"
    assert overviews[1].student_id == "00000003"


def test_excel_exporter_skips_demo_student(tmp_path):
    """DBにデモ生徒のデータが残っていてもExcel出力時にスキップされることを検証 (二重防御)"""
    db_file = tmp_path / "test_demo_skip.db"
    db_manager = DatabaseManager(str(db_file))
    repo = Repository(db_manager)

    # 1. ページ初期登録 (一般生徒とデモ生徒)
    ov1 = CurriculumOverview(
        classroom_name="テスト教室",
        classroom_code="999",
        school_year="2026",
        student_id="S001",
        student_name="山田 太郎",
        grade="高3",
        division="通常",
        subject="英語",
        row_index=0,
    )
    ov_demo = CurriculumOverview(
        classroom_name="テスト教室",
        classroom_code="999",
        school_year="2026",
        student_id="S002",
        student_name="デモ 一（デモ イチ）",
        grade="高2",
        division="通常",
        subject="数学",
        row_index=1,
    )
    repo.save_pages_init([ov1, ov_demo])

    # 2. 進捗サマリー登録
    prog1 = TargetProgress(
        classroom="テスト教室",
        student_id="S001",
        grade="高3",
        student_name="山田 太郎",
        division="通常",
        subject="英語",
        target_no="T1",
        target_name="1学期中間",
        start_date="2026/04/01",
        end_date="2026/05/20",
        period_str="2026/04/01～2026/05/20",
        is_current=True,
        textbooks_used="テキストA",
        target_score="80",
        current_score="85",
        target_unit_count=10,
        executed_count=8,
        executed_rate="80.0%",
        on_time_count=8,
        on_time_rate="100.0%",
        before_count=0,
        before_rate="0.0%",
        after_count=0,
        after_rate="0.0%",
        unexecuted_count=2,
        unexecuted_rate="20.0%",
        latest_execution_date="2026/05/10",
    )
    prog_demo = TargetProgress(
        classroom="テスト教室",
        student_id="S002",
        grade="高2",
        student_name="デモ 一（デモ イチ）",
        division="通常",
        subject="数学",
        target_no="T1",
        target_name="1学期中間",
        start_date="2026/04/01",
        end_date="2026/05/20",
        period_str="2026/04/01～2026/05/20",
        is_current=True,
        textbooks_used="テキストB",
        target_score="70",
        current_score="60",
        target_unit_count=5,
        executed_count=2,
        executed_rate="40.0%",
        on_time_count=2,
        on_time_rate="100.0%",
        before_count=0,
        before_rate="0.0%",
        after_count=0,
        after_rate="0.0%",
        unexecuted_count=3,
        unexecuted_rate="60.0%",
        latest_execution_date="2026/05/01",
    )

    eval1 = InstructionEvaluation(
        severity="PASS",
        author="講師A",
        student_info="真面目に取り組んでいます",
    )
    eval_demo = InstructionEvaluation(
        severity="PASS",
        author="講師B",
        student_info="テスト用生徒です",
    )

    detail1 = CurriculumDetailData(
        overview=ov1,
        raw_html="<html></html>",
        fetch_time="2026-10-28 10:00:00",
        targets=[],
        units=[],
        instruction_eval=eval1,
        progress_list=[prog1],
    )
    detail_demo = CurriculumDetailData(
        overview=ov_demo,
        raw_html="<html></html>",
        fetch_time="2026-10-28 10:00:00",
        targets=[],
        units=[],
        instruction_eval=eval_demo,
        progress_list=[prog_demo],
    )

    repo.save_parsed_result(detail1)
    repo.save_parsed_result(detail_demo)

    # ExcelExporter実行
    exporter = ExcelExporter(repo)
    out_xlsx = str(tmp_path / "demo_skip.xlsx")
    exporter.export_to(out_xlsx)
    wb = openpyxl.load_workbook(out_xlsx)

    # 各シートの生徒名を確認
    # 進捗率シート
    ws_prog = wb["進捗率"]
    prog_names = [ws_prog.cell(row=r, column=4).value for r in range(2, ws_prog.max_row + 1)]
    assert "山田 太郎" in prog_names
    assert "デモ 一（デモ イチ）" not in prog_names
    assert len(prog_names) == 1

    # 備考欄シート
    ws_inst = wb["備考欄"]
    inst_names = [ws_inst.cell(row=r, column=4).value for r in range(2, ws_inst.max_row + 1)]
    assert "山田 太郎" in inst_names
    assert "デモ 一（デモ イチ）" not in inst_names
    assert len(inst_names) == 1

    # サマリーシート (全体集計の生徒数が1であること)
    ws_sum = wb["サマリー"]
    # 2行目: 全体, 全件合計, 生徒数(col 3)
    total_students = ws_sum.cell(row=2, column=3).value
    assert total_students == 1


def test_copypaste_detector_skips_demo():
    """コピペ検出対象からデモ生徒が除外されることを検証"""
    eval_identical1 = InstructionEvaluation(
        severity="PASS",
        author="講師X",
        student_info="これは十分に長い生徒情報のテキストです。完全に同一の文面です。",
    )
    eval_identical2 = InstructionEvaluation(
        severity="PASS",
        author="講師X",
        student_info="これは十分に長い生徒情報のテキストです。完全に同一の文面です。",
    )
    eval_identical3 = InstructionEvaluation(
        severity="PASS",
        author="講師X",
        student_info="これは十分に長い生徒情報のテキストです。完全に同一の文面です。",
    )

    rec1 = CurriculumDetailData(
        overview=CurriculumOverview(
            classroom_name="校A", classroom_code="1", school_year="2026",
            student_id="S1", student_name="生徒1", grade="高1", division="通常", subject="英"
        ),
        raw_html="<html></html>",
        fetch_time="2026-10-28 10:00:00",
        targets=[], units=[], instruction_eval=eval_identical1
    )
    rec2 = CurriculumDetailData(
        overview=CurriculumOverview(
            classroom_name="校A", classroom_code="1", school_year="2026",
            student_id="S2", student_name="生徒2", grade="高1", division="通常", subject="英"
        ),
        raw_html="<html></html>",
        fetch_time="2026-10-28 10:00:00",
        targets=[], units=[], instruction_eval=eval_identical2
    )
    rec_demo = CurriculumDetailData(
        overview=CurriculumOverview(
            classroom_name="校A", classroom_code="1", school_year="2026",
            student_id="S3", student_name="デモ 一（デモ イチ）", grade="高1", division="通常", subject="英"
        ),
        raw_html="<html></html>",
        fetch_time="2026-10-28 10:00:00",
        targets=[], units=[], instruction_eval=eval_identical3
    )

    # 3件あれば通常はコピペ疑い警告が出るが、1件がデモ生徒なので有効件数は2件となり、警告は付与されない
    detect_copy_paste([rec1, rec2, rec_demo])
    assert not any("WARN_COPY_PASTE_SUSPECTED" in str(w) for w in rec1.instruction_eval.warnings)
    assert not any("WARN_COPY_PASTE_SUSPECTED" in str(w) for w in rec2.instruction_eval.warnings)
