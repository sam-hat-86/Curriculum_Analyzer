"""
Excel 5シート出力エンジン (仕様書§18, §20, §21, §24, §25, §26)
"""
import json
from datetime import datetime
from typing import Dict, List, Any
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from src.database.repository import Repository
from src.utils.constants import (
    SHEET_PROGRESS,
    SHEET_UNIT_DETAILS,
    SHEET_INSTRUCTION,
    SHEET_SUMMARY,
    SHEET_RAW,
    COLUMNS_PROGRESS,
    COLUMNS_UNIT_DETAILS,
    COLUMNS_INSTRUCTION,
    COLUMNS_SUMMARY,
    COLUMNS_RAW,
    DEFAULT_EMPTY_VALUE,
)

# スタイル定義
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid") # 濃紺
HEADER_FONT = Font(name="Yu Gothic UI", size=10, bold=True, color="FFFFFF")
DATA_FONT = Font(name="Yu Gothic UI", size=10)
BORDER_THIN = Side(border_style="thin", color="D9D9D9")
CELL_BORDER = Border(left=BORDER_THIN, right=BORDER_THIN, top=BORDER_THIN, bottom=BORDER_THIN)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)

class ExcelExporter:
    def __init__(self, repository: Repository):
        self.repo = repository

    def generate_workbook(self) -> openpyxl.Workbook:
        """データベースから全データを取得し、5シートを持つWorkbookを構築"""
        wb = openpyxl.Workbook()
        # デフォルトシートの名前変更
        ws_prog = wb.active
        ws_prog.title = SHEET_PROGRESS
        
        ws_unit = wb.create_sheet(title=SHEET_UNIT_DETAILS)
        ws_inst = wb.create_sheet(title=SHEET_INSTRUCTION)
        ws_sum = wb.create_sheet(title=SHEET_SUMMARY)
        ws_raw = wb.create_sheet(title=SHEET_RAW)

        # 1. 進捗率シート
        self._write_progress_sheet(ws_prog)

        # 2. 単元詳細シート
        self._write_unit_details_sheet(ws_unit)

        # 3. 備考欄シート
        self._write_instruction_sheet(ws_inst)

        # 4. サマリーシート
        self._write_summary_sheet(ws_sum)

        # 5. 原文シート
        self._write_raw_sheet(ws_raw)

        return wb

    def _apply_header_style(self, ws, columns: List[str]):
        """ヘッダー行の書式設定"""
        ws.append(columns)
        for col_idx in range(1, len(columns) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = ALIGN_CENTER
            cell.border = CELL_BORDER
        ws.freeze_panes = "A2"

    def _auto_column_widths(self, ws, max_len_cap: int = 40):
        """列幅の自動調整"""
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                # 改行がある場合は最長行
                lines = val.split("\n")
                line_max = max(len(l) for l in lines) if lines else 0
                max_len = max(max_len, line_max)
            adjusted_width = min(max(max_len + 3, 10), max_len_cap)
            ws.column_dimensions[col_letter].width = adjusted_width

    def _write_progress_sheet(self, ws):
        """[進捗率]シート出力 (全27列)"""
        self._apply_header_style(ws, COLUMNS_PROGRESS)
        rows = self.repo.get_all_progress_summaries()

        for r_idx, r in enumerate(rows, start=2):
            row_data = [
                r["classroom"] or DEFAULT_EMPTY_VALUE,
                r["student_id"] or DEFAULT_EMPTY_VALUE,
                r["grade"] or DEFAULT_EMPTY_VALUE,
                r["student_name"] or DEFAULT_EMPTY_VALUE,
                r["division"] or DEFAULT_EMPTY_VALUE,
                r["subject"] or DEFAULT_EMPTY_VALUE,
                r["target_no"] or DEFAULT_EMPTY_VALUE,
                r["target_name"] or DEFAULT_EMPTY_VALUE,
                r["start_date"] or DEFAULT_EMPTY_VALUE,
                r["end_date"] or DEFAULT_EMPTY_VALUE,
                r["period_str"] or DEFAULT_EMPTY_VALUE,
                bool(r["is_current"]), # Boolean
                r["textbooks_used"] or DEFAULT_EMPTY_VALUE,
                r["target_score"] or DEFAULT_EMPTY_VALUE,
                r["current_score"] or DEFAULT_EMPTY_VALUE,
                r["target_unit_count"],
                r["executed_count"],
                r["executed_rate"],
                r["on_time_count"],
                r["on_time_rate"],
                r["before_count"],
                r["before_rate"],
                r["after_count"],
                r["after_rate"],
                r["unexecuted_count"],
                r["unexecuted_rate"],
                r["latest_execution_date"] or DEFAULT_EMPTY_VALUE,
            ]
            ws.append(row_data)
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=r_idx, column=col_idx)
                cell.font = DATA_FONT
                cell.border = CELL_BORDER
                if col_idx in [12]: # 実施中 (Boolean)
                    cell.alignment = ALIGN_CENTER
                elif col_idx in [13]: # 使用教材 (改行あり)
                    cell.alignment = ALIGN_WRAP
                elif col_idx in [16, 17, 18, 19, 20, 21, 22, 23, 24, 25, 26]: # 数値・割合
                    cell.alignment = ALIGN_RIGHT
                else:
                    cell.alignment = ALIGN_LEFT

        self._auto_column_widths(ws)

    def _write_unit_details_sheet(self, ws):
        """[単元詳細]シート出力 (全23列)"""
        self._apply_header_style(ws, COLUMNS_UNIT_DETAILS)
        rows = self.repo.get_all_units()

        for r_idx, r in enumerate(rows, start=2):
            target_checks = {}
            if r["target_checks_json"]:
                try:
                    target_checks = json.loads(r["target_checks_json"])
                except Exception:
                    pass

            # ターゲットごとの展開 (1単元が複数ターゲットに該当する場合、それぞれ別行)
            # target_checks で True のターゲット、または単元単体
            active_targets = [t for t, checked in target_checks.items() if checked]
            if not active_targets:
                active_targets = ["-"]

            for t_no in active_targets:
                row_data = [
                    r["classroom_name"] or DEFAULT_EMPTY_VALUE,
                    r["student_id"] or DEFAULT_EMPTY_VALUE,
                    r["grade"] or DEFAULT_EMPTY_VALUE,
                    r["student_name"] or DEFAULT_EMPTY_VALUE,
                    r["division"] or DEFAULT_EMPTY_VALUE,
                    r["subject"] or DEFAULT_EMPTY_VALUE,
                    t_no,
                    DEFAULT_EMPTY_VALUE, # ターゲット名 (後で結合または参照)
                    DEFAULT_EMPTY_VALUE, # 開始日
                    DEFAULT_EMPTY_VALUE, # 終了日
                    r["unit_name"] or DEFAULT_EMPTY_VALUE,
                    r["middle_unit"] or DEFAULT_EMPTY_VALUE,
                    r["textbook"] or DEFAULT_EMPTY_VALUE,
                    r["problem_code"] or DEFAULT_EMPTY_VALUE,
                    r["execution_date"] or DEFAULT_EMPTY_VALUE,
                    r["timing"] or DEFAULT_EMPTY_VALUE,
                    bool(r["is_executed"]),
                    r["score"] or DEFAULT_EMPTY_VALUE,
                    r["max_score"] or DEFAULT_EMPTY_VALUE,
                    r["score_rate"] or DEFAULT_EMPTY_VALUE,
                    r["today_textbook"] or DEFAULT_EMPTY_VALUE,
                    r["execution_division"] or DEFAULT_EMPTY_VALUE,
                    r["confirm_test"] or DEFAULT_EMPTY_VALUE,
                ]
                ws.append(row_data)
                curr_row = ws.max_row
                for col_idx in range(1, len(row_data) + 1):
                    cell = ws.cell(row=curr_row, column=col_idx)
                    cell.font = DATA_FONT
                    cell.border = CELL_BORDER
                    if col_idx in [17]:
                        cell.alignment = ALIGN_CENTER
                    elif col_idx in [18, 19, 20]:
                        cell.alignment = ALIGN_RIGHT
                    else:
                        cell.alignment = ALIGN_LEFT

        self._auto_column_widths(ws)

    def _write_instruction_sheet(self, ws):
        """[備考欄]シート出力 (全18列)"""
        self._apply_header_style(ws, COLUMNS_INSTRUCTION)
        rows = self.repo.get_all_instruction_evaluations()

        for r_idx, r in enumerate(rows, start=2):
            fix_items = json.loads(r["fix_items_json"] or "[]")
            errors = json.loads(r["errors_json"] or "[]")
            warnings = json.loads(r["warnings_json"] or "[]")

            row_data = [
                r["classroom_name"] or DEFAULT_EMPTY_VALUE,
                r["student_id"] or DEFAULT_EMPTY_VALUE,
                r["grade"] or DEFAULT_EMPTY_VALUE,
                r["student_name"] or DEFAULT_EMPTY_VALUE,
                r["division"] or DEFAULT_EMPTY_VALUE,
                r["subject"] or DEFAULT_EMPTY_VALUE,
                r["severity"] or DEFAULT_EMPTY_VALUE,
                "\n".join(fix_items) if fix_items else DEFAULT_EMPTY_VALUE,
                "\n".join(errors) if errors else DEFAULT_EMPTY_VALUE,
                "\n".join(warnings) if warnings else DEFAULT_EMPTY_VALUE,
                r["author"] or DEFAULT_EMPTY_VALUE,
                r["target_school"] or DEFAULT_EMPTY_VALUE,
                r["textbook"] or DEFAULT_EMPTY_VALUE,
                r["plan"] or DEFAULT_EMPTY_VALUE,
                r["student_info"] or DEFAULT_EMPTY_VALUE,
                r["test"] or DEFAULT_EMPTY_VALUE,
                r["homework"] or DEFAULT_EMPTY_VALUE,
                r["unclassified_text"] or DEFAULT_EMPTY_VALUE,
            ]
            ws.append(row_data)
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=r_idx, column=col_idx)
                cell.font = DATA_FONT
                cell.border = CELL_BORDER
                if col_idx in [7]: # 判定
                    cell.alignment = ALIGN_CENTER
                elif col_idx in [8, 9, 10, 14, 15, 18]: # 複数行テキスト
                    cell.alignment = ALIGN_WRAP
                else:
                    cell.alignment = ALIGN_LEFT

        self._auto_column_widths(ws, max_len_cap=50)

    def _write_summary_sheet(self, ws):
        """[サマリー]シート出力 (仕様書§24: 全体、教室別、学年別、科目別集計)"""
        self._apply_header_style(ws, COLUMNS_SUMMARY)
        
        prog_rows = self.repo.get_all_progress_summaries()
        eval_rows = self.repo.get_all_instruction_evaluations()

        # グループ集計補助
        def summarize(group_key: str, key_name: str, records: List[Any], evals: List[Any]) -> List[Any]:
            student_ids = set()
            target_count = len(records)
            target_unit_count = 0
            executed_count = 0
            unexecuted_count = 0
            on_time_count = 0
            before_count = 0
            after_count = 0

            for r in records:
                student_ids.add(r["student_id"])
                target_unit_count += r["target_unit_count"] or 0
                executed_count += r["executed_count"] or 0
                unexecuted_count += r["unexecuted_count"] or 0
                on_time_count += r["on_time_count"] or 0
                before_count += r["before_count"] or 0
                after_count += r["after_count"] or 0

            error_cnt = 0
            warn_cnt = 0
            for e in evals:
                errs = json.loads(e["errors_json"] or "[]")
                warns = json.loads(e["warnings_json"] or "[]")
                error_cnt += len(errs)
                warn_cnt += len(warns)

            return [
                group_key,
                key_name,
                len(student_ids),
                target_count,
                target_unit_count,
                executed_count,
                unexecuted_count,
                on_time_count,
                before_count,
                after_count,
                error_cnt,
                warn_cnt,
            ]

        # 1. 全体集計
        summary_rows = []
        summary_rows.append(summarize("全体", "全件合計", prog_rows, eval_rows))

        # 2. 教室別集計
        classrooms = sorted(list(set(r["classroom"] for r in prog_rows if r["classroom"])))
        for cr in classrooms:
            cr_progs = [r for r in prog_rows if r["classroom"] == cr]
            cr_evals = [e for e in eval_rows if e["classroom_name"] == cr]
            summary_rows.append(summarize("教室別", cr, cr_progs, cr_evals))

        # 3. 学年別集計
        grades = sorted(list(set(r["grade"] for r in prog_rows if r["grade"])))
        for gr in grades:
            gr_progs = [r for r in prog_rows if r["grade"] == gr]
            gr_evals = [e for e in eval_rows if e["grade"] == gr]
            summary_rows.append(summarize("学年別", gr, gr_progs, gr_evals))

        # 4. 科目別集計
        subjects = sorted(list(set(r["subject"] for r in prog_rows if r["subject"])))
        for sub in subjects:
            sub_progs = [r for r in prog_rows if r["subject"] == sub]
            sub_evals = [e for e in eval_rows if e["subject"] == sub]
            summary_rows.append(summarize("科目別", sub, sub_progs, sub_evals))

        for r_idx, row_data in enumerate(summary_rows, start=2):
            ws.append(row_data)
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=r_idx, column=col_idx)
                cell.font = DATA_FONT
                cell.border = CELL_BORDER
                if col_idx in [1, 2]:
                    cell.alignment = ALIGN_LEFT
                else:
                    cell.alignment = ALIGN_RIGHT

        self._auto_column_widths(ws)

    def _write_raw_sheet(self, ws):
        """[原文]シート出力 (全6列)"""
        self._apply_header_style(ws, COLUMNS_RAW)
        rows = self.repo.get_all_raw_pages()

        for r_idx, r in enumerate(rows, start=2):
            html_snippet = r["html"] or ""
            # HTMLが巨大な場合はExcelセル制限(32767文字)を考慮し安全に格納
            if len(html_snippet) > 32000:
                html_snippet = html_snippet[:32000] + "... (truncated)"

            row_data = [
                r["url"] or DEFAULT_EMPTY_VALUE,
                r["fetch_time"] or str(r["created_at"]),
                "詳細ページ" if "detail" in (r["url"] or "") or "sim" in (r["url"] or "") else "授業データ",
                html_snippet or DEFAULT_EMPTY_VALUE,
                r["status"] or DEFAULT_EMPTY_VALUE,
                r["error_message"] or DEFAULT_EMPTY_VALUE,
            ]
            ws.append(row_data)
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=r_idx, column=col_idx)
                cell.font = DATA_FONT
                cell.border = CELL_BORDER
                if col_idx in [5]:
                    cell.alignment = ALIGN_CENTER
                else:
                    cell.alignment = ALIGN_LEFT

        self._auto_column_widths(ws, max_len_cap=50)
