"""
Excel 5シート出力および安全書き込みテスト (仕様書§18, §20, §21, §24, §25, §26, §28)
"""
import os
import pytest
import openpyxl
from src.database.db_manager import DatabaseManager
from src.database.repository import Repository
from src.exporter.excel_exporter import ExcelExporter
from src.exporter.safe_writer import SafeExcelWriter
from src.models.curriculum import (
    CurriculumOverview,
    TargetInfo,
    UnitRecord,
    TargetProgress,
    CurriculumDetailData,
)
from src.models.evaluation import InstructionEvaluation
from src.models.state import ProcessState, TimingCategory
from src.utils.constants import (
    SHEET_PROGRESS,
    SHEET_UNIT_DETAILS,
    SHEET_INSTRUCTION,
    SHEET_SUMMARY,
    SHEET_RAW,
    COLUMNS_PROGRESS,
    COLUMNS_UNIT_DETAILS,
    COLUMNS_INSTRUCTION,
    COLUMNS_SUMMARY,
    COLUMNS_RAW,
)

@pytest.fixture
def temp_db(tmp_path):
    db_file = str(tmp_path / "test_curriculum.db")
    db_manager = DatabaseManager(db_file)
    return db_manager

@pytest.fixture
def repo(temp_db):
    return Repository(temp_db)

def test_excel_five_sheets_export(repo, tmp_path):
    """5シート構成の生成テスト"""
    # モックデータの準備
    ov = CurriculumOverview(
        classroom_name="梅田本校",
        classroom_code="9901",
        school_year="2026",
        student_id="STU100",
        student_name="佐藤花子",
        grade="高1",
        division="通常",
        subject="英語",
        school_course="普通科",
        row_index=0,
        status=ProcessState.SAVED,
    )
    repo.save_pages_init([ov])

    target = TargetInfo(
        target_no="T1", target_name="1学期中間", start_date="2026/04/01", end_date="2026/05/31",
        period_str="2026/04/01～2026/05/31", is_current=True, target_score="90", current_score="85",
        site_progress_rate="100%",
    )

    unit = UnitRecord(
        unit_no="1", unit_name="Lesson 1", middle_unit="Part 1", textbook="Crown I",
        problem_code="E01", score="9", max_score="10", score_rate="90",
        execution_date="2026/04/15", timing=TimingCategory.ON_TIME, is_executed=True,
        today_textbook="Crown I", execution_division="通常", confirm_test="合格",
        target_checks={"T1": True}
    )

    progress = TargetProgress(
        classroom="梅田本校", student_id="STU100", grade="高1", student_name="佐藤花子",
        division="通常", subject="英語", target_no="T1", target_name="1学期中間",
        start_date="2026/04/01", end_date="2026/05/31", period_str="2026/04/01～2026/05/31",
        is_current=True, textbooks_used="Crown I", target_score="90", current_score="85",
        target_unit_count=1, executed_count=1, executed_rate="100", on_time_count=1,
        on_time_rate="100", before_count=0, before_rate="0", after_count=0, after_rate="0",
        unexecuted_count=0, unexecuted_rate="0", latest_execution_date="2026/04/15",
    )

    eval_res = InstructionEvaluation(
        severity="PASS", author="田中", target_school="公立高校", textbook="Crown I (所持)",
        plan="予定通り進める", student_info="意欲的", test="単語テスト", homework="ワークp.10"
    )

    detail_data = CurriculumDetailData(
        overview=ov,
        raw_html="<html><body>mock detail</body></html>",
        fetch_time="2026-09-18 10:00:00",
        raw_instruction="作成者：田中...",
        materials_listed=["Crown I"],
        targets=[target],
        units=[unit],
        instruction_eval=eval_res,
        progress_list=[progress],
    )
    repo.save_parsed_result(detail_data)

    # Excel出力テスト
    exporter = ExcelExporter(repo)
    safe_writer = SafeExcelWriter(exporter, str(tmp_path / "output"))

    # 1. 最終出力
    final_path = safe_writer.export_final()
    assert os.path.exists(final_path)

    # 2. シート検証
    wb = openpyxl.load_workbook(final_path)
    assert wb.sheetnames == [
        SHEET_PROGRESS,
        SHEET_UNIT_DETAILS,
        SHEET_INSTRUCTION,
        SHEET_SUMMARY,
        SHEET_RAW,
    ]

    # 進捗率シートの列数・データ確認
    ws_prog = wb[SHEET_PROGRESS]
    assert ws_prog.max_column == len(COLUMNS_PROGRESS)
    assert ws_prog.max_row == 2 # ヘッダー + 1行データ
    assert ws_prog.cell(row=2, column=1).value == "梅田本校"
    assert ws_prog.cell(row=2, column=4).value == "佐藤花子"
    assert ws_prog.cell(row=2, column=8).value == "1学期中間"
    assert ws_prog.cell(row=2, column=9).value == "OK" # 進捗遅れ列
    assert ws_prog.cell(row=2, column=17).value == 1 # 対象単元数

    # 単元詳細シートの列数・データ確認
    ws_unit = wb[SHEET_UNIT_DETAILS]
    assert ws_unit.max_column == len(COLUMNS_UNIT_DETAILS)
    assert ws_unit.max_row == 2
    assert ws_unit.cell(row=2, column=7).value == "T1"
    assert ws_unit.cell(row=2, column=8).value == "1学期中間"
    assert ws_unit.cell(row=2, column=9).value == "OK" # 進捗遅れ列
    assert ws_unit.cell(row=2, column=10).value == "2026/04/01"
    assert ws_unit.cell(row=2, column=11).value == "2026/05/31"
    assert ws_unit.cell(row=2, column=12).value == "Lesson 1"
    assert ws_unit.cell(row=2, column=17).value == "期間内"

    # 備考欄シートの確認
    ws_inst = wb[SHEET_INSTRUCTION]
    assert ws_inst.max_column == len(COLUMNS_INSTRUCTION)
    assert ws_inst.max_row == 2
    assert ws_inst.cell(row=2, column=7).value == "PASS"

    # サマリーシートの確認
    ws_sum = wb[SHEET_SUMMARY]
    assert ws_sum.max_column == len(COLUMNS_SUMMARY)
    assert ws_sum.max_row >= 2

    # 原文シートの確認
    ws_raw = wb[SHEET_RAW]
    assert ws_raw.max_column == len(COLUMNS_RAW)
    assert ws_raw.max_row == 2

    wb.close()

def test_clear_all_data(repo):
    """データ初期化機能の検証"""
    ov = CurriculumOverview(
        classroom_name="梅田本校", classroom_code="9901", school_year="2026",
        student_id="STU999", student_name="消去テスト", grade="高1", division="通常", subject="英語", row_index=0
    )
    repo.save_pages_init([ov])
    counts_before = repo.get_counts()
    assert counts_before["total"] >= 1

    repo.clear_all_data()
    counts_after = repo.get_counts()
    assert counts_after["total"] == 0
    assert counts_after["saved"] == 0
    assert counts_after["unfetched"] == 0


def test_excel_unit_details_multi_target_and_unassigned(repo, tmp_path):
    """単元の複数ターゲット展開および未割当単元の出力検証"""
    ov = CurriculumOverview(
        classroom_name="梅田本校", classroom_code="9901", school_year="2026",
        student_id="STU200", student_name="鈴木一郎", grade="高2", division="通常",
        subject="数学", school_course="理系", row_index=0, status=ProcessState.SAVED
    )
    repo.save_pages_init([ov])

    t1 = TargetInfo(
        target_no="T1", target_name="1学期中間", start_date="2026/05/01", end_date="2026/05/31",
        period_str="2026/05/01～2026/05/31", is_current=False, target_score="80", current_score="70",
        site_progress_rate="50%"
    )
    t2 = TargetInfo(
        target_no="T2", target_name="1学期期末", start_date="2026/06/15", end_date="2026/07/10",
        period_str="2026/06/15～2026/07/10", is_current=False, target_score="85", current_score="65",
        site_progress_rate="0%"
    )

    # 1. 複数ターゲットにチェックがある単元 (T1: 期間内, T2: 期間前)
    u_multi = UnitRecord(
        unit_no="1", unit_name="三角関数", middle_unit="加法定理", textbook="青チャート",
        problem_code="M01", score="10", max_score="10", score_rate="100",
        execution_date="2026/05/15", timing=TimingCategory.ON_TIME, is_executed=True,
        today_textbook="青チャート", execution_division="通常", confirm_test="合格",
        target_checks={"T1": True, "T2": True}
    )

    # 2. どのターゲットにもチェックがない単元 (未割当・実施済み)
    u_unassigned_exec = UnitRecord(
        unit_no="2", unit_name="複素数平面", middle_unit="極形式", textbook="青チャート",
        problem_code="M02", score="8", max_score="10", score_rate="80",
        execution_date="2026/05/20", timing=TimingCategory.ON_TIME, is_executed=True,
        today_textbook="青チャート", execution_division="通常", confirm_test="合格",
        target_checks={"T1": False, "T2": False}
    )

    # 3. どのターゲットにもチェックがない単元 (未割当・未実施)
    u_unassigned_unexec = UnitRecord(
        unit_no="3", unit_name="確率統計", middle_unit="条件付き確率", textbook="青チャート",
        problem_code="M03", score="-", max_score="-", score_rate="-",
        execution_date="-", timing=TimingCategory.UNEXECUTED, is_executed=False,
        today_textbook="-", execution_division="-", confirm_test="-",
        target_checks={"T1": False, "T2": False}
    )

    detail_data = CurriculumDetailData(
        overview=ov,
        raw_html="<html><body>test</body></html>",
        fetch_time="2026-09-18 10:00:00",
        raw_instruction="備考",
        materials_listed=["青チャート"],
        targets=[t1, t2],
        units=[u_multi, u_unassigned_exec, u_unassigned_unexec],
        instruction_eval=None,
        progress_list=[],
    )
    repo.save_parsed_result(detail_data)

    exporter = ExcelExporter(repo)
    safe_writer = SafeExcelWriter(exporter, str(tmp_path / "output_multi"))
    file_path = safe_writer.export_final()

    wb = openpyxl.load_workbook(file_path)
    ws = wb[SHEET_UNIT_DETAILS]

    # ヘッダー(行1) + u_multi(T1, T2の2行) + u_unassigned_exec(1行) + u_unassigned_unexec(1行) = 5行
    assert ws.max_row == 5

    # 行2: u_multi T1 (期間内: 2026/05/15 in 2026/05/01〜2026/05/31)
    assert ws.cell(row=2, column=7).value == "T1"
    assert ws.cell(row=2, column=8).value == "1学期中間"
    assert ws.cell(row=2, column=9).value == "OK" # 進捗遅れ列
    assert ws.cell(row=2, column=10).value == "2026/05/01"
    assert ws.cell(row=2, column=11).value == "2026/05/31"
    assert ws.cell(row=2, column=12).value == "三角関数"
    assert ws.cell(row=2, column=17).value == "期間内"

    # 行3: u_multi T2 (期間前: 2026/05/15 < 2026/06/15)
    assert ws.cell(row=3, column=7).value == "T2"
    assert ws.cell(row=3, column=8).value == "1学期期末"
    assert ws.cell(row=3, column=9).value == "OK" # 進捗遅れ列
    assert ws.cell(row=3, column=10).value == "2026/06/15"
    assert ws.cell(row=3, column=11).value == "2026/07/10"
    assert ws.cell(row=3, column=12).value == "三角関数"
    assert ws.cell(row=3, column=17).value == "期間前"

    # 行4: u_unassigned_exec (ターゲット未設定・実施済み)
    assert ws.cell(row=4, column=7).value == "-"
    assert ws.cell(row=4, column=8).value == "-"
    assert ws.cell(row=4, column=9).value == "OK" # 進捗遅れ列 (ターゲット未設定でもOK)
    assert ws.cell(row=4, column=10).value == "-"
    assert ws.cell(row=4, column=11).value == "-"
    assert ws.cell(row=4, column=12).value == "複素数平面"
    assert ws.cell(row=4, column=17).value == "期間内"

    # 行5: u_unassigned_unexec (ターゲット未設定・未実施)
    assert ws.cell(row=5, column=7).value == "-"
    assert ws.cell(row=5, column=8).value == "-"
    assert ws.cell(row=5, column=9).value == "OK" # 進捗遅れ列
    assert ws.cell(row=5, column=12).value == "確率統計"
    assert ws.cell(row=5, column=17).value == "未実施"

    wb.close()


def test_excel_export_progress_delay_column(repo, tmp_path):
    """進捗遅れのターゲットが存在する場合の進捗率シートおよび単元詳細シートの出力検証"""
    ov = CurriculumOverview(
        classroom_name="難波校",
        classroom_code="9902",
        school_year="2026",
        student_id="STU300",
        student_name="遅延確認生徒",
        grade="中3",
        division="通常",
        subject="数学",
        school_course="一般",
        row_index=0,
        status=ProcessState.SAVED,
    )
    repo.save_pages_init([ov])

    # 遅延ターゲット T1 と 正常ターゲット T2
    t1 = TargetInfo(
        target_no="T1",
        target_name="1学期中間/前期中間",
        progress_delay="進捗遅れ",
        start_date="2026/04/01",
        end_date="2026/05/20",
        period_str="2026/04/01～2026/05/20",
        is_current=True,
        target_score="80",
        current_score="60",
        site_progress_rate="40%",
    )
    t2 = TargetInfo(
        target_no="T2",
        target_name="1学期期末/前期中間2",
        progress_delay="OK",
        start_date="2026/06/01",
        end_date="2026/07/15",
        period_str="2026/06/01～2026/07/15",
        is_current=False,
        target_score="85",
        current_score="75",
        site_progress_rate="80%",
    )

    u1 = UnitRecord(
        unit_no="1", unit_name="平方根", middle_unit="計算", textbook="標準数学",
        problem_code="M10", score="7", max_score="10", score_rate="70",
        execution_date="2026/04/10", timing=TimingCategory.ON_TIME, is_executed=True,
        today_textbook="標準数学", execution_division="通常", confirm_test="合格",
        target_checks={"T1": True, "T2": False}
    )
    u2 = UnitRecord(
        unit_no="2", unit_name="二次方程式", middle_unit="解の公式", textbook="標準数学",
        problem_code="M11", score="8", max_score="10", score_rate="80",
        execution_date="2026/06/10", timing=TimingCategory.ON_TIME, is_executed=True,
        today_textbook="標準数学", execution_division="通常", confirm_test="合格",
        target_checks={"T1": False, "T2": True}
    )

    p1 = TargetProgress(
        classroom="難波校", student_id="STU300", grade="中3", student_name="遅延確認生徒",
        division="通常", subject="数学", target_no="T1", target_name="1学期中間/前期中間",
        progress_delay="進捗遅れ",
        start_date="2026/04/01", end_date="2026/05/20", period_str="2026/04/01～2026/05/20",
        is_current=True, textbooks_used="標準数学", target_score="80", current_score="60",
        target_unit_count=1, executed_count=1, executed_rate="100", on_time_count=1,
        on_time_rate="100", before_count=0, before_rate="0", after_count=0, after_rate="0",
        unexecuted_count=0, unexecuted_rate="0", latest_execution_date="2026/04/10",
    )
    p2 = TargetProgress(
        classroom="難波校", student_id="STU300", grade="中3", student_name="遅延確認生徒",
        division="通常", subject="数学", target_no="T2", target_name="1学期期末/前期中間2",
        progress_delay="OK",
        start_date="2026/06/01", end_date="2026/07/15", period_str="2026/06/01～2026/07/15",
        is_current=False, textbooks_used="標準数学", target_score="85", current_score="75",
        target_unit_count=1, executed_count=1, executed_rate="100", on_time_count=1,
        on_time_rate="100", before_count=0, before_rate="0", after_count=0, after_rate="0",
        unexecuted_count=0, unexecuted_rate="0", latest_execution_date="2026/06/10",
    )

    detail_data = CurriculumDetailData(
        overview=ov,
        raw_html="<html><body>mock detail</body></html>",
        fetch_time="2026-09-18 10:00:00",
        raw_instruction="",
        materials_listed=["標準数学"],
        targets=[t1, t2],
        units=[u1, u2],
        instruction_eval=None,
        progress_list=[p1, p2],
    )
    repo.save_parsed_result(detail_data)

    exporter = ExcelExporter(repo)
    safe_writer = SafeExcelWriter(exporter, str(tmp_path / "output_delay"))
    final_path = safe_writer.export_final()

    wb = openpyxl.load_workbook(final_path)

    # 1. 進捗率シートの検証
    ws_prog = wb[SHEET_PROGRESS]
    assert ws_prog.max_row == 3 # ヘッダー + 2レコード
    # 1行目データ (T1: 遅延あり)
    assert ws_prog.cell(row=2, column=8).value == "1学期中間/前期中間"
    assert "進捗遅れ" not in ws_prog.cell(row=2, column=8).value
    assert ws_prog.cell(row=2, column=9).value == "進捗遅れ"
    # 2行目データ (T2: 遅延なし)
    assert ws_prog.cell(row=3, column=8).value == "1学期期末/前期中間2"
    assert ws_prog.cell(row=3, column=9).value == "OK"

    # 2. 単元詳細シートの検証
    ws_unit = wb[SHEET_UNIT_DETAILS]
    assert ws_unit.max_row == 3 # ヘッダー + 2レコード
    # u1 (T1: 遅延あり)
    assert ws_unit.cell(row=2, column=8).value == "1学期中間/前期中間"
    assert ws_unit.cell(row=2, column=9).value == "進捗遅れ"
    # u2 (T2: 遅延なし)
    assert ws_unit.cell(row=3, column=8).value == "1学期期末/前期中間2"
    assert ws_unit.cell(row=3, column=9).value == "OK"

    wb.close()
