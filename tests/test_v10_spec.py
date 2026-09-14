"""
Curriculum Analyzer 仕様書 v10 検証テスト
- 生徒名抽出・正規化 (§2.1)
- 学年正規化 (§2.2)
- 学年を最優先とする志望校判定 (§3.1)
- 生徒名に基づくデモ除外 (§3.2)
- データモデル・キャッシュ復元 (§4)
- Excel出力各シートへの生徒名・学年出力 (§6)
"""

import os
from pathlib import Path
import pytest
import openpyxl

from core.models import AppConfig, CurriculumRecord, Severity
from core.normalization import extract_student_name_from_cell, normalize_grade
from core.evaluator import evaluate_single_record, check_record_exclusion
from core.cache import CacheManager
from core.excel_exporter import ExcelExporter
from core.constants import (
    ERR_SCHOOL_REQUIRED_HIGH3,
    ERR_SCHOOL_REQUIRED_JUNIOR3,
    ERR_SCHOOL_REQUIRED_ELEM6,
)


def test_extract_student_name_from_cell():
    """v10 §2.1, §5.1: 生徒名抽出および半角スペース正規化テスト"""
    # 1. 基本ケース: 学籍番号 + 改行 + 姓 + 半角空白 + 名
    assert extract_student_name_from_cell("00000001\n山田 太郎", "00000001") == "山田 太郎"

    # 2. 全角空白で区切られた姓名 -> 半角ASCIIスペース1個に結合
    assert extract_student_name_from_cell("00000001\n山田　太郎", "00000001") == "山田 太郎"

    # 3. 連続する空白や前後の空白
    assert extract_student_name_from_cell("00000001\n  山田   太郎  ", "00000001") == "山田 太郎"

    # 4. ボタンや操作系文字列の除外
    cell_with_buttons = "00000001\n山田 太郎\n詳細\n編集\n削除"
    assert extract_student_name_from_cell(cell_with_buttons, "00000001") == "山田 太郎"

    # 5. 空白なしの単一氏名
    assert extract_student_name_from_cell("00000001\n山田太郎", "00000001") == "山田太郎"

    # 6. 学籍番号のみ（氏名なし）
    assert extract_student_name_from_cell("00000001", "00000001") == ""
    assert extract_student_name_from_cell("00000001\n", "00000001") == ""

    # 7. 空文字またはNone
    assert extract_student_name_from_cell("", "00000001") == ""
    assert extract_student_name_from_cell(None, "00000001") == ""


def test_normalize_grade():
    """v10 §2.2, §5.2: 学年正規化テスト (NFKC、全角数字を半角化、学年推測なし)"""
    # 全角数字 -> 半角数字
    assert normalize_grade("高２") == "高2"
    assert normalize_grade("中３") == "中3"
    assert normalize_grade("小６") == "小6"

    # 複数学年
    assert normalize_grade("高１・２") == "高1・2"
    assert normalize_grade("高1・2") == "高1・2"

    # 前後空白除去
    assert normalize_grade("  中2  ") == "中2"

    # 空値
    assert normalize_grade("") == ""
    assert normalize_grade(None) == ""


def test_grade_driven_school_requirement():
    """v10 §3.1: 学年を最優先正式インプットとする志望校判定テスト"""
    config = AppConfig(base_url="https://example.com")
    instruction_no_school = """作成者:山田
教材:フォレスタ(所持)
進め方:演習を進めます
生徒情報:真面目
小テスト:単語
宿題:p.10"""

    instruction_with_school = """作成者:山田
志望校:東京大学
教材:フォレスタ(所持)
進め方:演習を進めます
生徒情報:真面目
小テスト:単語
宿題:p.10"""

    # 高3: 志望校必須 (未記載なら ERR_SCHOOL_REQUIRED_HIGH3)
    rec_high3 = CurriculumRecord(
        student_id="STU_H3",
        division="通常授業",
        subject="英語",
        grade="高3",
        raw_instruction=instruction_no_school,
    )
    res_high3 = evaluate_single_record(rec_high3, config)
    assert ERR_SCHOOL_REQUIRED_HIGH3 in res_high3.errors

    rec_high3_ok = CurriculumRecord(
        student_id="STU_H3",
        division="通常授業",
        subject="英語",
        grade="高3",
        raw_instruction=instruction_with_school,
    )
    res_high3_ok = evaluate_single_record(rec_high3_ok, config)
    assert ERR_SCHOOL_REQUIRED_HIGH3 not in res_high3_ok.errors
    assert res_high3_ok.severity == Severity.PASS

    # 中3: 志望校必須 (未記載なら ERR_SCHOOL_REQUIRED_JUNIOR3)
    rec_mid3 = CurriculumRecord(
        student_id="STU_M3",
        division="通常授業",
        subject="数学",
        grade="中3",
        raw_instruction=instruction_no_school,
    )
    res_mid3 = evaluate_single_record(rec_mid3, config)
    assert ERR_SCHOOL_REQUIRED_JUNIOR3 in res_mid3.errors

    # 高2: 志望校任意 (未記載でもエラーにならない)
    rec_high2 = CurriculumRecord(
        student_id="STU_H2",
        division="通常授業",
        subject="英語",
        grade="高2",
        raw_instruction=instruction_no_school,
    )
    res_high2 = evaluate_single_record(rec_high2, config)
    assert ERR_SCHOOL_REQUIRED_HIGH3 not in res_high2.errors
    assert ERR_SCHOOL_REQUIRED_JUNIOR3 not in res_high2.errors
    assert res_high2.severity == Severity.PASS

    # 中1: 志望校任意
    rec_mid1 = CurriculumRecord(
        student_id="STU_M1",
        division="通常授業",
        subject="数学",
        grade="中1",
        raw_instruction=instruction_no_school,
    )
    res_mid1 = evaluate_single_record(rec_mid1, config)
    assert res_mid1.severity == Severity.PASS


def test_demo_exclusion_by_student_name():
    """v10 §3.2: 生徒名（姓がデモ）に基づく除外判定テスト"""
    # 姓が「デモ」 -> 除外 ("DEMO")
    rec_demo1 = CurriculumRecord(student_id="STU_001", student_name="デモ 太郎", division="通常授業", subject="英語", raw_instruction="test")
    assert check_record_exclusion(rec_demo1) == "DEMO"

    rec_demo2 = CurriculumRecord(student_id="STU_002", student_name="デモ", division="通常授業", subject="数学", raw_instruction="test")
    assert check_record_exclusion(rec_demo2) == "DEMO"

    # 通常生徒名 -> 除外されない (None)
    rec_normal = CurriculumRecord(student_id="STU_003", student_name="山田 太郎", division="通常授業", subject="国語", raw_instruction="test")
    assert check_record_exclusion(rec_normal) is None

    # 「デモ田」など名字がデモそのものでない場合 -> 除外されない (None)
    rec_demoda = CurriculumRecord(student_id="STU_004", student_name="デモ田 次郎", division="通常授業", subject="国語", raw_instruction="test")
    assert check_record_exclusion(rec_demoda) is None


def test_v10_cache_student_name_and_grade(tmp_path):
    """v10 §4: キャッシュ (JSONL) への生徒名・学年の保存と完全復元テスト"""
    mgr = CacheManager(str(tmp_path))
    records = [
        CurriculumRecord(
            student_id="STU_1001",
            student_name="山田 太郎",
            grade="高3",
            division="通常授業",
            subject="英語",
            raw_instruction="作成者:講師A",
            classroom_name="天王寺校",
            classroom_code="C01",
            school_year="2026年度",
        ),
        CurriculumRecord(
            student_id="STU_1002",
            student_name="佐藤 花子",
            grade="中2",
            division="夏期講習",
            subject="数学",
            raw_instruction="作成者:講師B",
            classroom_name="梅田校",
            classroom_code="C02",
            school_year="2026年度",
        ),
    ]

    mgr.save(records, "1.0.0")
    assert mgr.exists()

    loaded, skip_count, status = mgr.load()
    assert status == "ok"
    assert skip_count == 0
    assert len(loaded) == 2

    assert loaded[0].student_id == "STU_1001"
    assert loaded[0].student_name == "山田 太郎"
    assert loaded[0].grade == "高3"
    assert loaded[0].classroom_name == "天王寺校"

    assert loaded[1].student_id == "STU_1002"
    assert loaded[1].student_name == "佐藤 花子"
    assert loaded[1].grade == "中2"


def test_v10_excel_export_sheets_and_student_columns(tmp_path):
    """v10 §6: Excel出力全対象シートに生徒名・学年が正しく配置されることの検証"""
    config = AppConfig(base_url="https://example.com")
    rec = CurriculumRecord(
        student_id="STU_8888",
        student_name="田中 一郎",
        grade="高3",
        division="通常授業",
        subject="英語",
        raw_instruction="""作成者:山田
志望校:京都大学
教材:体系数学(所持)
進め方:第1回から基本問題の演習を進めます
生徒情報:真面目で学習態度は大変良好です
小テスト:英単語
宿題:p.15""",
        classroom_name="京都校",
        classroom_code="K01",
        school_year="2026年度",
    )
    res = evaluate_single_record(rec, config)

    exporter = ExcelExporter()
    out_file = exporter.export(str(tmp_path), [rec], [res], timestamp="20260914_v10test")
    assert os.path.exists(out_file)

    wb = openpyxl.load_workbook(out_file)

    # 1. チェック結果 (20列)
    # [教室名, 教室コード, 年度, 学籍番号, 生徒名, 学年, 受講区分, 科目, 判定, ...]
    ws_check = wb["チェック結果"]
    assert ws_check.cell(row=2, column=1).value == "京都校"
    assert ws_check.cell(row=2, column=2).value == "K01"
    assert ws_check.cell(row=2, column=3).value == "2026年度"
    assert ws_check.cell(row=2, column=4).value == "STU_8888"
    assert ws_check.cell(row=2, column=5).value == "田中 一郎"
    assert ws_check.cell(row=2, column=6).value == "高3"
    assert ws_check.cell(row=2, column=7).value == "通常授業"
    assert ws_check.cell(row=2, column=8).value == "英語"
    assert ws_check.cell(row=2, column=9).value == "PASS"

    # 2. 原文 (10列)
    # [教室名, 教室コード, 年度, 学籍番号, 生徒名, 学年, 受講区分, 科目, 判定, 原文備考欄]
    ws_raw = wb["原文"]
    assert ws_raw.cell(row=2, column=1).value == "京都校"
    assert ws_raw.cell(row=2, column=2).value == "K01"
    assert ws_raw.cell(row=2, column=3).value == "2026年度"
    assert ws_raw.cell(row=2, column=4).value == "STU_8888"
    assert ws_raw.cell(row=2, column=5).value == "田中 一郎"
    assert ws_raw.cell(row=2, column=6).value == "高3"
    assert ws_raw.cell(row=2, column=7).value == "通常授業"
    assert ws_raw.cell(row=2, column=8).value == "英語"
    assert ws_raw.cell(row=2, column=9).value == "PASS"

    # 3. 解析詳細 (10列)
    # [学籍番号, 生徒名, 学年, 行番号, 原文行, 検出見出し, 正規化見出し, 割当セクション, 信頼度, 本文]
    ws_trace = wb["解析詳細"]
    assert ws_trace.cell(row=2, column=1).value == "STU_8888"
    assert ws_trace.cell(row=2, column=2).value == "田中 一郎"
    assert ws_trace.cell(row=2, column=3).value == "高3"

    # 4. 教材詳細 (13列)
    # [学籍番号, 生徒名, 学年, 受講区分, 科目, 教室名, 教室コード, 年度, 教材, ステータス, ステータス分類, 判定, 補足]
    ws_tb = wb["教材詳細"]
    assert ws_tb.cell(row=2, column=1).value == "STU_8888"
    assert ws_tb.cell(row=2, column=2).value == "田中 一郎"
    assert ws_tb.cell(row=2, column=3).value == "高3"
    assert ws_tb.cell(row=2, column=4).value == "通常授業"
    assert ws_tb.cell(row=2, column=5).value == "英語"
    assert ws_tb.cell(row=2, column=6).value == "京都校"
    assert ws_tb.cell(row=2, column=7).value == "K01"
    assert ws_tb.cell(row=2, column=8).value == "2026年度"
    assert ws_tb.cell(row=2, column=9).value == "体系数学"
    assert ws_tb.cell(row=2, column=10).value == "所持"
