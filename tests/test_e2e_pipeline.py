"""
エンドツーエンド (E2E) パイプライン統合テスト
- sampleMSL.html の一覧解析
- me.html, late.html の詳細パース
- 単一DBWriterによるSQLite保存
- SafeExcelWriterによる5シート出力
"""
import os
import queue
import time
import pytest
import openpyxl

from src.database.db_manager import DatabaseManager
from src.database.repository import Repository
from src.database.db_writer import DBWriter
from src.parser.list_parser import ListParser
from src.parser.detail_parser import DetailParser
from src.analyzer.instruction_evaluator import InstructionEvaluator
from src.analyzer.progress_calculator import calculate_target_progress
from src.exporter.excel_exporter import ExcelExporter
from src.exporter.safe_writer import SafeExcelWriter

def test_e2e_pipeline(tmp_path):
    db_file = str(tmp_path / "e2e_test.db")
    out_dir = str(tmp_path / "output_v2")

    db_manager = DatabaseManager(db_file)
    repo = Repository(db_manager)

    # 1. 一覧パース
    list_parser = ListParser()
    with open("v1/sampleMSL.html", "r", encoding="utf-8", errors="ignore") as f:
        list_html = f.read()
    overviews = list_parser.parse(list_html, classroom_name="天王寺本校", classroom_code="43210", school_year="2026年度")
    assert len(overviews) == 11

    repo.save_pages_init(overviews)

    # 2. DBWriter起動
    save_queue = queue.Queue()
    db_writer = DBWriter(repository=repo, save_queue=save_queue)
    db_writer.start()

    # 3. 詳細パース (me.html と late.html を投入)
    detail_parser = DetailParser()
    evaluator = InstructionEvaluator()

    with open("v1/me.html", "r", encoding="utf-8", errors="ignore") as f:
        me_html = f.read()
    with open("v1/late.html", "r", encoding="utf-8", errors="ignore") as f:
        late_html = f.read()

    # 1件目 (me.html)
    d1 = detail_parser.parse(me_html, overview=overviews[0])
    d1.instruction_eval = evaluator.evaluate(d1.raw_instruction, grade=d1.overview.grade)
    d1.progress_list = [calculate_target_progress(d1.overview, t, d1.units) for t in d1.targets]
    save_queue.put(d1)

    # 2件目 (late.html)
    d2 = detail_parser.parse(late_html, overview=overviews[1])
    d2.instruction_eval = evaluator.evaluate(d2.raw_instruction, grade=d2.overview.grade)
    d2.progress_list = [calculate_target_progress(d2.overview, t, d2.units) for t in d2.targets]
    save_queue.put(d2)

    # 待機・DBWriter終了
    save_queue.join()
    db_writer.stop()
    save_queue.put(None)
    db_writer.join()

    # 4. DB確認
    counts = repo.get_counts()
    assert counts["saved"] == 2

    # 5. Excel 5シート出力
    exporter = ExcelExporter(repo)
    safe_writer = SafeExcelWriter(exporter, output_dir=out_dir)
    excel_path = safe_writer.export_final()

    assert os.path.exists(excel_path)

    wb = openpyxl.load_workbook(excel_path)
    assert len(wb.sheetnames) == 5

    # 進捗率シート: 1件目(T1=1行) + 2件目(T1〜T5=5行) = 計6行データ (+ ヘッダー1行 = 7行)
    ws_prog = wb["進捗率"]
    assert ws_prog.max_row == 7
    assert ws_prog.max_column == 27

    # 単元詳細シート
    ws_unit = wb["単元詳細"]
    assert ws_unit.max_row > 50
    assert ws_unit.max_column == 23

    # 備考欄シート: 2行データ (+ ヘッダー1行 = 3行)
    ws_inst = wb["備考欄"]
    assert ws_inst.max_row == 3
    assert ws_inst.max_column == 18

    # サマリーシート
    ws_sum = wb["サマリー"]
    assert ws_sum.max_row >= 2
    assert ws_sum.max_column == 12

    wb.close()
